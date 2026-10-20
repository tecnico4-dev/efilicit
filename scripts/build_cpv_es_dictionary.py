#!/usr/bin/env python3
"""Regenera apps/web/src/lib/data/cpv-es.json desde el XLSX oficial TED/SIMAP.

Fuente: https://ted.europa.eu/documents/d/ted/cpv_2008_xls
Uso:
  python scripts/build_cpv_es_dictionary.py [ruta/cpv_2008_ver_2013.xlsx]
Si no se pasa ruta, descarga el ZIP oficial a un temporal.
"""
from __future__ import annotations

import json
import re
import sys
import tempfile
import zipfile
from datetime import date
from pathlib import Path
from urllib.request import urlretrieve

SOURCE_URL = "https://ted.europa.eu/documents/d/ted/cpv_2008_xls"
REPO_ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = REPO_ROOT / "apps" / "web" / "src" / "lib" / "data" / "cpv-es.json"


def _load_workbook(xlsx_path: Path):
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise SystemExit(
            "Necesitas openpyxl: pip install openpyxl"
        ) from exc
    return openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)


def download_xlsx(dest_dir: Path) -> Path:
    zip_path = dest_dir / "cpv_2008_xls.zip"
    print(f"Downloading {SOURCE_URL} …")
    urlretrieve(SOURCE_URL, zip_path)
    with zipfile.ZipFile(zip_path) as zf:
        names = [n for n in zf.namelist() if n.lower().endswith(".xlsx")]
        if not names:
            raise SystemExit("ZIP sin .xlsx")
        zf.extract(names[0], dest_dir)
        return dest_dir / names[0]


def find_parent(code: str, code_set: set[str]) -> str | None:
    for length in range(7, 1, -1):
        cand = code[:length].ljust(8, "0")
        if cand != code and cand in code_set:
            return cand
    return None


def build(xlsx_path: Path) -> dict:
    wb = _load_workbook(xlsx_path)
    ws = wb["CPV codes"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    es_i = header.index("ES")
    code_i = header.index("CODE")

    codes: dict[str, dict] = {}
    for row in rows[1:]:
        raw = str(row[code_i] or "").strip()
        label = str(row[es_i] or "").strip()
        if not raw or not label:
            continue
        m = re.match(r"^(\d{8})(?:-(\d))?$", raw)
        if not m:
            continue
        code8, check = m.group(1), m.group(2)
        node: dict = {
            "label": label,
            "parent": None,
            "children": [],
        }
        if check is not None:
            node["check"] = check
        codes[code8] = node
    wb.close()

    code_set = set(codes)
    for code in codes:
        parent = find_parent(code, code_set)
        codes[code]["parent"] = parent
        if parent is not None:
            codes[parent]["children"].append(code)
    for node in codes.values():
        node["children"].sort()

    return {
        "meta": {
            "vocabulary": "CPV",
            "version": "2008",
            "language": "es",
            "source": "TED/SIMAP cpv_2008_ver_2013.xlsx (hoja CPV codes, columna ES)",
            "source_url": SOURCE_URL,
            "generated_at": date.today().isoformat(),
            "code_count": len(codes),
        },
        "codes": dict(sorted(codes.items())),
    }


def main() -> None:
    if len(sys.argv) > 1:
        xlsx = Path(sys.argv[1])
        data = build(xlsx)
    else:
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = download_xlsx(Path(tmp))
            data = build(xlsx)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(
        json.dumps(data, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    sample = data["codes"]["33600000"]
    print(f"Wrote {OUT_PATH} ({data['meta']['code_count']} codes)")
    print(f"33600000 -> {sample['label']} (parent={sample['parent']})")


if __name__ == "__main__":
    main()
